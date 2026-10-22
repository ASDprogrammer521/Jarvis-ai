from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from actions._common import log, need, open_path, output_dir, safe_name, unique_path

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
    _XL = True
except ImportError:
    _XL = False

THEMES = {
    "blue":  dict(head="1F3A5F", accent="2E75B6", band="EAF1F8", title="1F3A5F"),
    "green": dict(head="1E5B3A", accent="2E8B57", band="EAF5EE", title="1E5B3A"),
    "dark":  dict(head="23283B", accent="6EA8FF", band="F0F2F7", title="23283B"),
    "mono":  dict(head="333333", accent="777777", band="F2F2F2", title="111111"),
    "sunset": dict(head="8A2E1F", accent="E4572E", band="FBEDE8", title="8A2E1F"),
}
FORMATS = {
    "currency": '"$"#,##0.00', "percent": "0.0%", "number": "#,##0.00", "integer": "#,##0",
    "date": "yyyy-mm-dd", "text": "@", "general": "General",
}
FUNCS = {"sum": "SUM", "avg": "AVERAGE", "average": "AVERAGE", "min": "MIN", "max": "MAX", "count": "COUNT"}
_NUM = re.compile(r"^-?\d+(?:\.\d+)?$")


def _col(ref) -> int:
    if isinstance(ref, int):
        return ref
    s = str(ref).strip()
    return int(s) if s.isdigit() else column_index_from_string(s.upper())


def _coerce(v):
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        if _NUM.match(s):
            return float(s) if "." in s else int(s)
    return v


def _build_sheet(ws, sh: dict, th: dict) -> None:
    cols = sh.get("columns") or []
    rows = sh.get("rows") or []
    ncols = max([len(cols)] + [len(r) for r in rows]) if (cols or rows) else 1
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    r = 1
    if sh.get("title"):
        ws.cell(r, 1, sh["title"]).font = Font(size=16, bold=True, color=th["title"])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, ncols))
        ws.row_dimensions[r].height = 26
        r += 1
        if sh.get("subtitle"):
            ws.cell(r, 1, sh["subtitle"]).font = Font(size=11, italic=True, color="666666")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, ncols))
            r += 1
        r += 1
    header_row = r
    fmts: dict[int, str] = {}
    if cols:
        for i, c in enumerate(cols, 1):
            h = c.get("header", "") if isinstance(c, dict) else str(c)
            cell = ws.cell(r, i, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=th["head"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            if isinstance(c, dict):
                fmts[i] = FORMATS.get(str(c.get("format", "")).lower(), c.get("format") or "")
        ws.row_dimensions[r].height = 22
        r += 1
    first = r
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row, 1):
            v = _coerce(v)
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{row}", str(r))
            cell = ws.cell(r, ci, v)
            cell.border = border
            if fmts.get(ci):
                cell.number_format = fmts[ci]
            if isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("=")):
                cell.alignment = Alignment(horizontal="right")
            if sh.get("zebra", True) and ri % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=th["band"])
        r += 1
    last = r - 1
    tot = sh.get("totals")
    if tot and rows:
        label_col = _col(tot.get("label_column", 1))
        ws.cell(r, label_col, tot.get("label", "Total")).font = Font(bold=True)
        for ref, fn in (tot.get("columns") or {}).items():
            ci = _col(ref)
            L = get_column_letter(ci)
            f = FUNCS.get(str(fn).lower(), "SUM")
            cell = ws.cell(r, ci, f"={f}({L}{first}:{L}{last})")
            cell.number_format = fmts.get(ci, "General")
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="right")
        for ci in range(1, ncols + 1):
            c = ws.cell(r, ci)
            c.fill = PatternFill("solid", fgColor=th["band"])
            c.border = Border(top=Side(style="medium", color=th["head"]), bottom=thin, left=thin, right=thin)
            if c.font is None or not c.font.bold:
                c.font = Font(bold=True)
        r += 1
    for i in range(1, ncols + 1):
        L = get_column_letter(i)
        want = None
        if i - 1 < len(cols) and isinstance(cols[i - 1], dict):
            want = cols[i - 1].get("width")
        if want:
            ws.column_dimensions[L].width = float(want)
        else:
            longest = max([len(str(ws.cell(rr, i).value or "")) for rr in range(header_row, min(r, header_row + 60))] + [8])
            ws.column_dimensions[L].width = min(48, longest + 3)
    if cols and sh.get("freeze", True):
        ws.freeze_panes = ws.cell(first, 1)
    if cols and sh.get("autofilter", True) and last >= first:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last}"
    for rule in sh.get("conditional") or []:
        try:
            L = get_column_letter(_col(rule["column"]))
            op = {"gt": "greaterThan", "lt": "lessThan", "ge": "greaterThanOrEqual", "le": "lessThanOrEqual",
                  "eq": "equal", "ne": "notEqual"}.get(str(rule.get("rule", "gt")), "greaterThan")
            color = str(rule.get("color", "C6EFCE")).lstrip("#")
            ws.conditional_formatting.add(
                f"{L}{first}:{L}{last}",
                CellIsRule(operator=op, formula=[str(rule.get("value", 0))], fill=PatternFill("solid", bgColor=color, fgColor=color)))
        except Exception:
            pass
    for ch in sh.get("charts") or []:
        try:
            kind = str(ch.get("type", "bar")).lower()
            chart = {"line": LineChart, "pie": PieChart}.get(kind, BarChart)()
            chart.title = ch.get("title")
            chart.height, chart.width = 8.5, 16
            cat_i = _col(ch.get("cat_col", 1))
            for dc in ch.get("data_cols") or []:
                di = _col(dc)
                chart.add_data(Reference(ws, min_col=di, min_row=header_row, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=cat_i, min_row=first, max_row=last))
            ws.add_chart(chart, ch.get("anchor") or f"{get_column_letter(ncols + 2)}{header_row}")
        except Exception as e:
            print(f"[Excel] chart skipped: {e}")
    ws.page_setup.orientation = "landscape" if ncols > 6 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.sheet_view.showGridLines = bool(sh.get("gridlines", False))


def build_workbook(spec: dict, path: Path) -> int:
    th = THEMES.get(str(spec.get("theme", "blue")).lower(), THEMES["blue"])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = spec.get("sheets") or []
    if not sheets:
        raise ValueError("the spreadsheet has no sheets")
    used = set()
    for i, sh in enumerate(sheets, 1):
        name = re.sub(r"[\\/*?:\[\]]", "", str(sh.get("name") or f"Sheet{i}"))[:31] or f"Sheet{i}"
        while name in used:
            name = (name[:28] + f"_{i}")
        used.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = th["accent"]
        _build_sheet(ws, sh, th)
    wb.save(str(path))
    return len(sheets)


def _spec_from_csv(csv_path: Path, title: str) -> dict:
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        data = list(csv.reader(fh))
    if not data:
        raise ValueError("the CSV is empty")
    head, body = data[0], data[1:]
    return {"sheets": [{"name": csv_path.stem[:31], "title": title or csv_path.stem,
                        "columns": [{"header": h} for h in head], "rows": body}]}


def build_spreadsheet(parameters: dict, player=None, speak=None) -> str:
    if not _XL:
        return need("openpyxl")
    p = parameters or {}
    try:
        spec = None
        if p.get("spec_json"):
            spec = json.loads(p["spec_json"]) if isinstance(p["spec_json"], str) else p["spec_json"]
        elif p.get("csv_path"):
            spec = _spec_from_csv(Path(str(p["csv_path"])).expanduser(), str(p.get("title") or ""))
        if not isinstance(spec, dict):
            return "Sir, I need either the spreadsheet content (spec_json) or a CSV file path."
        if p.get("theme"):
            spec["theme"] = p["theme"]
        title = str(p.get("title") or spec.get("title") or (spec.get("sheets") or [{}])[0].get("title") or "spreadsheet")
        path = unique_path(output_dir("Spreadsheets"), safe_name(p.get("filename") or title, "spreadsheet"), ".xlsx")
        n = build_workbook(spec, path)
        log(f"[Excel] Saved {path}", player)
        if p.get("open_after", True):
            open_path(path)
        return f"Spreadsheet ready, sir: {n} sheet(s) with formatting and formulas. Saved to {path}."
    except Exception as e:
        return f"Sir, building the spreadsheet failed: {e}"


TOOL = {
    "name": "build_spreadsheet",
    "description": (
        "Builds a formatted Excel (.xlsx) workbook from structured content or a CSV file: styled headers, number "
        "formats, formulas (cells starting with '=', use {row} for the current row), total rows, charts, "
        "conditional formatting, frozen headers and filters. Use when the user already has the data/structure; "
        "to design a spreadsheet from a description use generate_spreadsheet instead."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "spec_json": {"type": "STRING", "description": (
                "JSON: {theme, sheets:[{name,title,subtitle,columns:[{header,format,width}],rows:[[...]],"
                "totals:{label,columns:{\"C\":\"sum\"}},charts:[{type,title,data_cols:[\"C\"],cat_col:\"A\"}],"
                "conditional:[{column,rule,value,color}]}]}")},
            "csv_path": {"type": "STRING", "description": "Path of a CSV file to turn into a formatted workbook"},
            "title": {"type": "STRING", "description": "Report title"},
            "theme": {"type": "STRING", "description": "blue | green | dark | mono | sunset"},
            "filename": {"type": "STRING", "description": "Optional file name"},
            "open_after": {"type": "BOOLEAN", "description": "Open the file when done (default true)"},
        },
    },
    "handler": build_spreadsheet,
}

# by ASDcuber521
