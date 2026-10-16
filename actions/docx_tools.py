from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from actions._common import ask_json, log, need, open_path, output_dir, safe_name, unique_path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    _DOCX = True
except ImportError:
    _DOCX = False

THEMES = {
    "modern":  dict(font="Calibri", head="1F3A5F", accent="2E75B6", band="DCE9F5", size=11),
    "classic": dict(font="Times New Roman", head="000000", accent="444444", band="EEEEEE", size=12),
    "minimal": dict(font="Arial", head="222222", accent="888888", band="F3F3F3", size=10.5),
    "bold":    dict(font="Segoe UI", head="7A1F3D", accent="C2185B", band="F8E1EA", size=11),
}

SYSTEM = (
    "You write polished Word documents. Reply with ONE JSON object and nothing else: "
    '{"title": str, "subtitle": str, "blocks": [block, ...]}. A block is one of: '
    '{"type":"heading","level":1|2|3,"text":str} | {"type":"paragraph","text":str} | '
    '{"type":"bullets","items":[str]} | {"type":"numbered","items":[str]} | '
    '{"type":"table","headers":[str],"rows":[[str]],"caption":str} | {"type":"quote","text":str,"author":str} | '
    '{"type":"page_break"}. Text may use **bold** and *italic* markers. '
    "Structure the document logically with an introduction and conclusion, real content (no placeholders), "
    "and write in the requested language."
)


def _rgb(h: str):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(style_or_run, name: str, size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = _rgb(color)
    el = style_or_run.element if hasattr(style_or_run, "element") else None
    try:
        rpr = (style_or_run._element.get_or_add_rPr() if hasattr(style_or_run, "_element") and hasattr(style_or_run._element, "get_or_add_rPr")
               else style_or_run.element.get_or_add_rPr())
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(a), name)
    except Exception:
        pass


def _shade(cell, hex_fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _add_field(paragraph, instr: str):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _rich(paragraph, text: str, theme: dict, size=None, color=None, italic=False):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", str(text))
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**") and len(part) > 4
        ital = (not bold) and part.startswith("*") and part.endswith("*") and len(part) > 2
        run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if ital else part)
        run.bold = bold or None
        run.italic = (ital or italic) or None
        _set_font(run, theme["font"], size, color=color)


def setup_document(doc, theme: dict):
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    _set_font(normal, theme["font"], theme["size"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for lvl, sz in ((1, 20), (2, 15), (3, 12.5)):
        st = doc.styles[f"Heading {lvl}"]
        _set_font(st, theme["font"], sz, True, theme["head"] if lvl < 3 else theme["accent"])
        st.paragraph_format.space_before = Pt(16 if lvl == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        _set_font(doc.styles[name], theme["font"], theme["size"])
        doc.styles[name].paragraph_format.space_after = Pt(3)
    if "Footer" in [x.name for x in doc.styles]:
        _set_font(doc.styles["Footer"], theme["font"], 9, color="888888")
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run("Page ")
    _set_font(r, theme["font"], 9, color="888888")
    _add_field(foot, "PAGE")
    for run in foot.runs:
        _set_font(run, theme["font"], 9, color="888888")


def add_title(doc, title: str, subtitle: str, theme: dict):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    _set_font(r, theme["font"], 28, True, theme["head"])
    if subtitle:
        s = doc.add_paragraph()
        r = s.add_run(subtitle)
        _set_font(r, theme["font"], 14, color=theme["accent"])
    d = doc.add_paragraph()
    r = d.add_run(datetime.now().strftime("%d %B %Y"))
    _set_font(r, theme["font"], 9.5, color="888888")
    rule = doc.add_paragraph()
    ppr = rule._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", theme["accent"])):
        bt.set(qn(k), v)
    bd.append(bt)
    ppr.append(bd)


def add_table(doc, block: dict, theme: dict):
    headers = [str(h) for h in block.get("headers") or []]
    rows = [[str(c) for c in r] for r in block.get("rows") or []]
    cols = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
    if not cols:
        return
    t = doc.add_table(rows=(1 if headers else 0) + len(rows), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    ri = 0
    if headers:
        for i in range(cols):
            c = t.cell(0, i)
            c.text = ""
            run = c.paragraphs[0].add_run(headers[i] if i < len(headers) else "")
            _set_font(run, theme["font"], theme["size"] - 0.5, True, "FFFFFF")
            _shade(c, theme["head"])
        ri = 1
    for r_i, row in enumerate(rows):
        for i in range(cols):
            c = t.cell(ri + r_i, i)
            c.text = ""
            run = c.paragraphs[0].add_run(row[i] if i < len(row) else "")
            _set_font(run, theme["font"], theme["size"] - 0.5)
            if r_i % 2 == 1:
                _shade(c, theme["band"])
    if block.get("caption"):
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(str(block["caption"]))
        _set_font(r, theme["font"], 9.5, color="666666")
        r.italic = True
    else:
        doc.add_paragraph()


def render_blocks(doc, blocks: list, theme: dict) -> int:
    n = 0
    for b in blocks or []:
        if not isinstance(b, dict):
            continue
        t = str(b.get("type", "paragraph")).lower()
        if t == "heading":
            h = doc.add_heading(level=max(1, min(3, int(b.get("level") or 1))))
            _rich(h, b.get("text", ""), theme)
            for r in h.runs:
                _set_font(r, theme["font"], None, True)
        elif t == "paragraph":
            _rich(doc.add_paragraph(), b.get("text", ""), theme)
        elif t in ("bullets", "numbered"):
            style = "List Bullet" if t == "bullets" else "List Number"
            for it in b.get("items") or []:
                _rich(doc.add_paragraph(style=style), it, theme)
        elif t == "table":
            add_table(doc, b, theme)
        elif t == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            _rich(p, "“" + str(b.get("text", "")) + "”", theme, size=theme["size"] + 1, color=theme["accent"], italic=True)
            if b.get("author"):
                a = doc.add_paragraph()
                a.paragraph_format.left_indent = Cm(1)
                r = a.add_run("— " + str(b["author"]))
                _set_font(r, theme["font"], theme["size"] - 1, color="777777")
        elif t == "page_break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "image" and b.get("path") and Path(str(b["path"])).exists():
            doc.add_picture(str(b["path"]), width=Inches(float(b.get("width") or 5.5)))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            continue
        n += 1
    return n


def _plan(p: dict, player):
    raw = p.get("content_json") or p.get("blocks")
    if raw:
        try:
            data = json.loads(raw) if isinstance(raw, str) else raw
            if isinstance(data, list):
                data = {"blocks": data}
            if isinstance(data, dict) and data.get("blocks"):
                return data
        except Exception:
            pass
    topic = str(p.get("topic") or p.get("instructions") or "").strip()
    if not topic:
        return None
    log(f"[Word] Drafting: {topic[:80]}", player)
    lang = p.get("language") or "the same language as the request"
    return ask_json(f"Write a document about: {topic}\nLength: {p.get('length') or 'about 2 pages'}\nLanguage: {lang}\n"
                    f"Tone: {p.get('tone') or 'professional'}", SYSTEM, default=None)


def _replace_in_paragraph(par, old: str, new: str) -> int:
    if old not in par.text:
        return 0
    cnt = par.text.count(old)
    for run in par.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return cnt
    full = par.text.replace(old, new)
    for i, run in enumerate(par.runs):
        run.text = full if i == 0 else ""
    return cnt


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            yield p
        for p in sec.footer.paragraphs:
            yield p


def word_document(parameters: dict, player=None, speak=None) -> str:
    if not _DOCX:
        return need("docx", "python-docx")
    p = parameters or {}
    action = str(p.get("action") or "create").lower().strip()
    theme = THEMES.get(str(p.get("style") or "modern").lower(), THEMES["modern"])

    try:
        if action == "create":
            plan = _plan(p, player)
            if not plan:
                return "Sir, tell me what the document should be about, or give me its content."
            doc = Document()
            setup_document(doc, theme)
            title = str(plan.get("title") or p.get("title") or p.get("topic") or "Document")
            add_title(doc, title, str(plan.get("subtitle") or ""), theme)
            n = render_blocks(doc, plan.get("blocks"), theme)
            path = unique_path(output_dir("Documents"), safe_name(p.get("filename") or title, "document"), ".docx")
            doc.save(str(path))
            log(f"[Word] Saved {path}", player)
            if p.get("open_after", True):
                open_path(path)
            return f"Document ready, sir: '{title}' with {n} sections. Saved to {path}."

        path = Path(str(p.get("path") or "")).expanduser()
        if not path.exists() or path.suffix.lower() != ".docx":
            return "Sir, I need the path of an existing .docx file."

        if action == "read":
            doc = Document(str(path))
            out = []
            for par in doc.paragraphs:
                if not par.text.strip():
                    continue
                s = par.style.name if par.style is not None else ""
                out.append(("# " * int(s[-1]) if s.startswith("Heading ") and s[-1].isdigit() else "") + par.text)
            text = "\n".join(out)
            return text[:6000] + ("\n…(truncated)" if len(text) > 6000 else "") or "The document is empty."

        if action == "append":
            plan = _plan(p, player)
            if not plan:
                return "Sir, tell me what to add to the document."
            doc = Document(str(path))
            n = render_blocks(doc, plan.get("blocks"), theme)
            doc.save(str(path))
            return f"Added {n} sections to {path.name}, sir."

        if action == "replace":
            old, new = str(p.get("find") or ""), str(p.get("replace_with") or "")
            if not old:
                return "Sir, what text should I find?"
            doc = Document(str(path))
            total = sum(_replace_in_paragraph(par, old, new) for par in _all_paragraphs(doc))
            doc.save(str(path))
            return f"Replaced {total} occurrence(s) of '{old}' in {path.name}, sir."

        return "Unknown action. Available: create, append, replace, read."
    except Exception as e:
        return f"Sir, the Word operation failed: {e}"


TOOL = {
    "name": "word_document",
    "description": (
        "Creates and edits Word (.docx) documents with real typography: title block, headings, paragraphs, "
        "bullet and numbered lists, formatted tables, quotes, page numbers. Actions: create (from a topic or "
        "structured content), append (add sections to an existing file), replace (find and replace text), "
        "read (return the text). Use for reports, letters, essays, CVs, notes, contracts."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {"type": "STRING", "description": "create | append | replace | read (default create)"},
            "topic": {"type": "STRING", "description": "What to write about (create/append)"},
            "title": {"type": "STRING", "description": "Document title"},
            "content_json": {"type": "STRING", "description": "Optional ready-made content: JSON {title, subtitle, blocks:[...]}"},
            "path": {"type": "STRING", "description": "Existing .docx path (append/replace/read)"},
            "find": {"type": "STRING", "description": "Text to find (replace)"},
            "replace_with": {"type": "STRING", "description": "Replacement text (replace)"},
            "style": {"type": "STRING", "description": "modern | classic | minimal | bold"},
            "length": {"type": "STRING", "description": "e.g. '1 page', '3 pages'"},
            "tone": {"type": "STRING", "description": "e.g. formal, friendly"},
            "language": {"type": "STRING", "description": "Language to write in"},
            "filename": {"type": "STRING", "description": "Optional file name"},
            "open_after": {"type": "BOOLEAN", "description": "Open the file when done (default true)"},
        },
    },
    "handler": word_document,
}

# by ASDcuber521
