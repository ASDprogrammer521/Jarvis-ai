from __future__ import annotations

import json
import platform
import shutil
import subprocess
import tempfile
from pathlib import Path

from actions._common import cfg, log, need, open_path, output_dir, safe_name, unique_path

try:
    from pypdf import PdfReader, PdfWriter
    _PYPDF = True
except ImportError:
    _PYPDF = False

_OS = platform.system()
_IMG_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
_OFFICE_EXT = {".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp", ".xlsx", ".xls", ".ods", ".csv", ".html", ".htm", ".txt", ".md"}


def _paths(p: dict) -> list[Path]:
    raw = p.get("paths") or p.get("path") or p.get("files") or []
    if isinstance(raw, str):
        try:
            data = json.loads(raw)
            raw = data if isinstance(data, list) else [raw]
        except Exception:
            raw = [x for x in raw.replace(";", "\n").replace(",", "\n").split("\n")]
    return [Path(str(x).strip().strip('"')).expanduser() for x in raw if str(x).strip()]


def _parse_pages(spec, total: int) -> list[int]:
    if not spec:
        return list(range(total))
    out: list[int] = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, _, b = part.partition("-")
            lo = int(a) if a else 1
            hi = int(b) if b else total
            out += [i - 1 for i in range(max(1, lo), min(total, hi) + 1)]
        else:
            n = int(part)
            if 1 <= n <= total:
                out.append(n - 1)
    return out


def _ranges(spec: str, total: int) -> list[list[int]]:
    groups = []
    for chunk in str(spec).replace(" ", "").split(";" if ";" in str(spec) else ","):
        g = _parse_pages(chunk, total)
        if g:
            groups.append(g)
    return groups


def _soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(c).exists():
            return c
    return None


def _ocr_page_images(path: Path, pages: list[int], lang: str):
    try:
        import pytesseract
    except ImportError:
        raise RuntimeError(need("pytesseract") + " (and install the Tesseract OCR program).")
    custom = cfg().get("tesseract_path")
    if custom:
        pytesseract.pytesseract.tesseract_cmd = custom
    elif _OS == "Windows" and not shutil.which("tesseract"):
        for c in (r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if Path(c).exists():
                pytesseract.pytesseract.tesseract_cmd = c
                break
    images = []
    try:
        import fitz
        doc = fitz.open(str(path))
        for i in pages:
            pix = doc[i].get_pixmap(dpi=220)
            from PIL import Image
            images.append((i, Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
    except ImportError:
        try:
            from pdf2image import convert_from_path
        except ImportError:
            raise RuntimeError(need("pymupdf") + " (recommended) so scanned pages can be read.")
        for i in pages:
            imgs = convert_from_path(str(path), dpi=220, first_page=i + 1, last_page=i + 1)
            if imgs:
                images.append((i, imgs[0]))
    out = {}
    for i, img in images:
        out[i] = pytesseract.image_to_string(img, lang=lang).strip()
    return out


def _extract_text(path: Path, pages_spec, ocr: str, lang: str) -> tuple[str, int, int]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise RuntimeError("the PDF is password-protected")
    total = len(reader.pages)
    idx = _parse_pages(pages_spec, total)
    texts: dict[int, str] = {}
    for i in idx:
        try:
            texts[i] = (reader.pages[i].extract_text() or "").strip()
        except Exception:
            texts[i] = ""
    ocr_pages = 0
    if ocr != "never":
        need_ocr = [i for i in idx if ocr == "always" or len(texts[i]) < 20]
        if need_ocr:
            res = _ocr_page_images(path, need_ocr, lang)
            for i, t in res.items():
                if t:
                    texts[i] = t
                    ocr_pages += 1
    body = "\n\n".join(f"--- Page {i + 1} ---\n{texts[i]}" for i in idx)
    return body, total, ocr_pages


def _convert(src: Path, folder: Path) -> Path:
    ext = src.suffix.lower()
    out = unique_path(folder, safe_name(src.stem, "converted"), ".pdf")
    if ext == ".pdf":
        raise RuntimeError("that file is already a PDF")
    if ext in _IMG_EXT:
        from PIL import Image
        img = Image.open(src)
        frames = []
        try:
            n = getattr(img, "n_frames", 1)
            for k in range(n):
                img.seek(k)
                frames.append(img.convert("RGB").copy())
        except Exception:
            frames = [img.convert("RGB")]
        frames[0].save(str(out), "PDF", resolution=150.0, save_all=True, append_images=frames[1:])
        return out
    if ext in _OFFICE_EXT:
        so = _soffice()
        if so:
            with tempfile.TemporaryDirectory() as td:
                proc = subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
                                      capture_output=True, text=True, timeout=180)
                made = list(Path(td).glob("*.pdf"))
                if made:
                    shutil.move(str(made[0]), str(out))
                    return out
                raise RuntimeError(f"LibreOffice could not convert it ({proc.stderr.strip()[:120]})")
        if ext == ".docx" and _OS in ("Windows", "Darwin"):
            try:
                from docx2pdf import convert
                convert(str(src), str(out))
                if out.exists():
                    return out
            except ImportError:
                raise RuntimeError("install LibreOffice (free) or run: pip install docx2pdf (needs Microsoft Word)")
        if ext in (".txt", ".md"):
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
                from xml.sax.saxutils import escape
                styles = getSampleStyleSheet()
                story = []
                for line in src.read_text(encoding="utf-8", errors="replace").splitlines():
                    story.append(Paragraph(escape(line) or "&nbsp;", styles["BodyText"]))
                    story.append(Spacer(1, 2))
                SimpleDocTemplate(str(out), pagesize=A4).build(story)
                return out
            except ImportError:
                pass
        raise RuntimeError("install LibreOffice (free, libreoffice.org) so I can convert Office documents")
    raise RuntimeError(f"I can't convert '{ext}' files to PDF")


def _tables(path: Path, pages_spec, folder: Path) -> tuple[str, int]:
    try:
        import pdfplumber
    except ImportError:
        raise RuntimeError(need("pdfplumber"))
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    count, preview = 0, ""
    with pdfplumber.open(str(path)) as pdf:
        for i in _parse_pages(pages_spec, len(pdf.pages)):
            for t_i, table in enumerate(pdf.pages[i].extract_tables() or []):
                rows = [[(c or "").strip() for c in r] for r in table if r and any((c or "").strip() for c in r)]
                if not rows:
                    continue
                count += 1
                ws = wb.create_sheet(f"P{i + 1}_T{t_i + 1}"[:31])
                for r in rows:
                    ws.append(r)
                for c in ws[1]:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="1F3A5F")
                    c.alignment = Alignment(horizontal="center")
                for col in ws.columns:
                    ws.column_dimensions[col[0].column_letter].width = min(45, max(10, max(len(str(c.value or "")) for c in col) + 2))
                if not preview:
                    preview = "\n".join(" | ".join(r) for r in rows[:6])
    if not count:
        return "", 0
    out = unique_path(folder, safe_name(path.stem, "tables") + "_tables", ".xlsx")
    wb.save(str(out))
    return f"{out}\n{preview}", count


def pdf_tools(parameters: dict, player=None, speak=None) -> str:
    if not _PYPDF:
        return need("pypdf")
    p = parameters or {}
    action = str(p.get("action") or "").lower().strip()
    folder = output_dir("PDF")
    files = _paths(p)

    try:
        if action == "merge":
            if len(files) < 2:
                return "Sir, give me at least two PDF files to merge."
            w = PdfWriter()
            for f in files:
                if not f.exists():
                    return f"Sir, I can't find {f}."
                w.append(PdfReader(str(f)))
            out = unique_path(folder, safe_name(p.get("filename") or "merged"), ".pdf")
            with open(out, "wb") as fh:
                w.write(fh)
            log(f"[PDF] Merged {len(files)} files → {out}", player)
            if p.get("open_after", False):
                open_path(out)
            return f"Merged {len(files)} PDFs into one, {len(w.pages)} pages. Saved to {out}."

        if not files:
            return "Sir, which file should I use?"
        src = files[0]
        if not src.exists():
            return f"Sir, I can't find {src}."

        if action == "split":
            reader = PdfReader(str(src))
            total = len(reader.pages)
            if p.get("ranges"):
                groups = _ranges(p["ranges"], total)
            elif p.get("every"):
                n = max(1, int(p["every"]))
                groups = [list(range(i, min(i + n, total))) for i in range(0, total, n)]
            else:
                groups = [[i] for i in range(total)]
            if not groups:
                return "Sir, those page ranges don't match this document."
            sub = folder / safe_name(src.stem + "_split")
            sub.mkdir(parents=True, exist_ok=True)
            made = []
            for k, g in enumerate(groups, 1):
                w = PdfWriter()
                for i in g:
                    w.add_page(reader.pages[i])
                label = f"p{g[0] + 1}" if len(g) == 1 else f"p{g[0] + 1}-{g[-1] + 1}"
                o = sub / f"{safe_name(src.stem)}_{k:02d}_{label}.pdf"
                with open(o, "wb") as fh:
                    w.write(fh)
                made.append(o)
            log(f"[PDF] Split into {len(made)} files in {sub}", player)
            if p.get("open_after", False):
                open_path(sub)
            return f"Split {src.name} ({total} pages) into {len(made)} files. Saved in {sub}."

        if action in ("extract_text", "text", "ocr"):
            ocr = "always" if action == "ocr" else str(p.get("ocr") or "auto").lower()
            if ocr not in ("auto", "always", "never"):
                ocr = "auto"
            body, total, used = _extract_text(src, p.get("pages"), ocr, str(p.get("language") or "eng"))
            if not body.strip():
                return "I found no readable text in that PDF, sir."
            out = unique_path(folder, safe_name(src.stem) + "_text", ".txt")
            out.write_text(body, encoding="utf-8")
            note = f" (OCR used on {used} page(s))" if used else ""
            head = body[:1800] + ("\n…" if len(body) > 1800 else "")
            return f"Extracted text from {src.name}{note}. Full text saved to {out}.\n\n{head}"

        if action in ("extract_tables", "tables"):
            res, n = _tables(src, p.get("pages"), folder)
            if not n:
                return "I found no tables in that PDF, sir."
            xlsx, _, prev = res.partition("\n")
            if p.get("open_after", False):
                open_path(xlsx)
            return f"Found {n} table(s) and saved them to an Excel file: {xlsx}\nFirst table preview:\n{prev}"

        if action == "convert":
            out = _convert(src, folder)
            log(f"[PDF] Converted {src.name} → {out}", player)
            if p.get("open_after", False):
                open_path(out)
            return f"Converted {src.name} to PDF. Saved to {out}."

        if action == "info":
            r = PdfReader(str(src))
            meta = r.metadata or {}
            first = (r.pages[0].extract_text() or "").strip()[:200] if len(r.pages) else ""
            return (f"{src.name}: {len(r.pages)} pages, {src.stat().st_size / 1024:.0f} KB, "
                    f"encrypted: {r.is_encrypted}, title: {meta.get('/Title') or '-'}, author: {meta.get('/Author') or '-'}."
                    f"{(' Starts with: ' + first) if first else ' No selectable text (likely scanned).'}")

        return "Unknown action. Available: merge, split, extract_text, ocr, extract_tables, convert, info."
    except Exception as e:
        return f"Sir, the PDF operation failed: {e}"


TOOL = {
    "name": "pdf_tools",
    "description": (
        "PDF toolkit. Actions: merge (combine several PDFs), split (by page ranges, every N pages, or one file per "
        "page), extract_text (with automatic OCR for scanned pages), ocr (force OCR), extract_tables (tables to an "
        "Excel file), convert (Word/PowerPoint/Excel/text/images to PDF), info (page count and metadata). "
        "Use whenever the user works with PDF files."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {"type": "STRING", "description": "merge | split | extract_text | ocr | extract_tables | convert | info"},
            "paths": {"type": "ARRAY", "items": {"type": "STRING"}, "description": "File path(s). merge needs 2+, other actions use the first"},
            "pages": {"type": "STRING", "description": "Page selection like '1-3,5,8-' (extract_text, extract_tables)"},
            "ranges": {"type": "STRING", "description": "split: groups like '1-3;4-6;7-'"},
            "every": {"type": "INTEGER", "description": "split: pages per output file"},
            "ocr": {"type": "STRING", "description": "auto | always | never (extract_text)"},
            "language": {"type": "STRING", "description": "OCR language code(s), e.g. 'eng', 'eng+rus'"},
            "filename": {"type": "STRING", "description": "Output name for merge"},
            "open_after": {"type": "BOOLEAN", "description": "Open the result when done"},
        },
        "required": ["action"],
    },
    "handler": pdf_tools,
}

# by ASDcuber521
