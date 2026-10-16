from __future__ import annotations

import json
import math
import re
from pathlib import Path

from actions._common import ask_json, log, need, open_path, output_dir, safe_name, unique_path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    _PPTX = True
except ImportError:
    _PPTX = False

PALETTES = {
    "midnight": dict(bg="0B1020", surface="151C36", text="F2F5FF", muted="9AA6D0", accent="6EE7FF", accent2="B388FF"),
    "ocean":    dict(bg="F4F8FB", surface="FFFFFF", text="0F2A3F", muted="5B7388", accent="0E7FBF", accent2="20C7B2"),
    "sunset":   dict(bg="1B1020", surface="2A1730", text="FFF3EC", muted="C9A7B8", accent="FF7A59", accent2="FFC857"),
    "forest":   dict(bg="F3F7F2", surface="FFFFFF", text="1C3326", muted="5E7A68", accent="2E8B57", accent2="9BC53D"),
    "mono":     dict(bg="FFFFFF", surface="F2F2F2", text="111111", muted="666666", accent="111111", accent2="888888"),
    "aurora":   dict(bg="070B14", surface="101A2E", text="EAF6FF", muted="8FB0C8", accent="5CFFB0", accent2="6EA8FF"),
}

SW, SH = 13.333, 7.5
FONT = "Calibri"
SYSTEM = (
    "You design professional slide decks. Reply with ONE JSON object and nothing else. Schema: "
    '{"title": str, "subtitle": str, "slides": [{"layout": one of '
    '"bullets"|"two_column"|"comparison"|"process"|"timeline"|"stats"|"quote"|"section", '
    '"title": str, "bullets": [str], "left": {"heading": str, "bullets": [str]}, '
    '"right": {"heading": str, "bullets": [str]}, "steps": [{"title": str, "text": str}], '
    '"stats": [{"value": str, "label": str}], "quote": str, "author": str, "notes": str}]}. '
    "Rules: concise bullets (max 14 words, 3-5 per slide); every slide has speaker notes of 2-4 sentences; "
    "vary layouts and use at least one diagram layout (process/timeline) when it fits; "
    "do not include the title slide or the closing slide, they are added automatically; "
    "write in the requested language."
)


def _c(h: str):
    return RGBColor.from_string(h.upper().lstrip("#"))


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = _c(fill)
    if line:
        s.line.color.rgb = _c(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size=18, color="000000", bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = FONT
        r.font.color.rgb = _c(color)
    return tb


def _lines(text: str, size: float, width_in: float) -> int:
    chars_per_line = max(8, int(width_in * 72 / (size * 0.52)))
    return max(1, math.ceil(len(text) / chars_per_line))


def _bullets(slide, x, y, w, h, items, pal, size=20):
    items = [str(i) for i in (items or []) if str(i).strip()]
    if not items:
        return
    total = sum(_lines(i, size, w - 0.45) for i in items)
    while size > 12 and (total * size * 1.32 + len(items) * 20) / 72 > h:
        size -= 1
        total = sum(_lines(i, size, w - 0.45) for i in items)
    cy = y
    for it in items:
        n = _lines(it, size, w - 0.45)
        rh = (n * size * 1.32 + 20) / 72
        _rect(slide, x + 0.02, cy + (size * 0.55) / 72, 0.12, 0.12, pal["accent"], MSO_SHAPE.OVAL)
        _text(slide, x + 0.3, cy, w - 0.3, rh, it, size, pal["text"])
        cy += rh


def _base(prs, pal, title=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _c(pal["bg"])
    if title:
        _rect(slide, 0.6, 0.62, 0.09, 0.62, pal["accent"])
        size = 34 if len(title) <= 38 else 28 if len(title) <= 60 else 24
        _text(slide, 0.85, 0.5, SW - 1.7, 0.9, title, size, pal["text"], True, anchor=MSO_ANCHOR.MIDDLE)
    if number is not None:
        _text(slide, SW - 1.3, SH - 0.6, 0.8, 0.35, str(number), 11, pal["muted"], align=PP_ALIGN.RIGHT)
    return slide


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = str(text)


def _title_slide(prs, pal, title, subtitle, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    f = slide.background.fill
    f.gradient()
    f.gradient_angle = 35
    f.gradient_stops[0].color.rgb = _c(pal["bg"])
    f.gradient_stops[1].color.rgb = _c(pal["surface"])
    _rect(slide, 0.9, 2.55, 1.6, 0.09, pal["accent"])
    _rect(slide, 2.6, 2.55, 0.5, 0.09, pal["accent2"])
    size = 54 if len(title) <= 28 else 44 if len(title) <= 48 else 36
    _text(slide, 0.85, 2.75, SW - 2.4, 1.9, title, size, pal["text"], True, anchor=MSO_ANCHOR.TOP)
    if subtitle:
        _text(slide, 0.9, 4.75, SW - 3, 1.0, subtitle, 22, pal["muted"])
    _rect(slide, SW - 3.2, -1.4, 4.2, 4.2, pal["accent2"], MSO_SHAPE.OVAL)
    _rect(slide, SW - 2.2, 5.2, 3.6, 3.6, pal["accent"], MSO_SHAPE.OVAL)
    _notes(slide, notes)


def _closing_slide(prs, pal, text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    f = slide.background.fill
    f.gradient()
    f.gradient_angle = 215
    f.gradient_stops[0].color.rgb = _c(pal["bg"])
    f.gradient_stops[1].color.rgb = _c(pal["surface"])
    _rect(slide, SW / 2 - 0.8, 2.7, 1.6, 0.09, pal["accent"])
    _text(slide, 1, 2.95, SW - 2, 1.3, text, 48, pal["text"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def _slide_bullets(prs, pal, s, n):
    slide = _base(prs, pal, s.get("title", ""), n)
    _bullets(slide, 0.9, 1.8, SW - 1.8, 4.9, s.get("bullets"), pal, 30)
    return slide


def _col(slide, pal, x, y, w, h, block):
    _rect(slide, x, y, w, h, pal["surface"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    _rect(slide, x, y, w, 0.09, pal["accent"])
    head = (block or {}).get("heading", "")
    if head:
        _text(slide, x + 0.3, y + 0.25, w - 0.6, 0.7, head, 26, pal["accent"], True)
    _bullets(slide, x + 0.3, y + 1.1, w - 0.6, h - 1.3, (block or {}).get("bullets"), pal, 24)


def _slide_two(prs, pal, s, n):
    slide = _base(prs, pal, s.get("title", ""), n)
    w = (SW - 1.2 - 0.5) / 2
    _col(slide, pal, 0.6, 1.8, w, 4.8, s.get("left"))
    _col(slide, pal, 0.6 + w + 0.5, 1.8, w, 4.8, s.get("right"))
    return slide


def _slide_compare(prs, pal, s, n):
    slide = _slide_two(prs, pal, s, n)
    _rect(slide, SW / 2 - 0.35, 3.9, 0.7, 0.7, pal["accent2"], MSO_SHAPE.OVAL)
    _text(slide, SW / 2 - 0.35, 3.9, 0.7, 0.7, "VS", 16, pal["bg"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return slide


def _slide_process(prs, pal, s, n):
    slide = _base(prs, pal, s.get("title", ""), n)
    steps = (s.get("steps") or [])[:6]
    k = max(1, len(steps))
    gap = 0.35
    w = (SW - 1.2 - gap * (k - 1)) / k
    for i, st in enumerate(steps):
        x = 0.6 + i * (w + gap)
        _rect(slide, x, 2.4, w, 3.4, pal["surface"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        _rect(slide, x + w / 2 - 0.35, 2.05, 0.7, 0.7, pal["accent"], MSO_SHAPE.OVAL)
        _text(slide, x + w / 2 - 0.35, 2.05, 0.7, 0.7, str(i + 1), 20, pal["bg"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        _text(slide, x + 0.15, 2.95, w - 0.3, 0.8, st.get("title", ""), 22 if k <= 4 else 17, pal["text"], True, PP_ALIGN.CENTER)
        _text(slide, x + 0.15, 3.75, w - 0.3, 1.9, st.get("text", ""), 17 if k <= 4 else 14, pal["muted"], align=PP_ALIGN.CENTER)
        if i < k - 1:
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w + 0.03), Inches(4.1),
                                           Inches(x + w + gap - 0.03), Inches(4.1))
            c.line.color.rgb = _c(pal["accent2"])
            c.line.width = Pt(2.5)
    return slide


def _slide_timeline(prs, pal, s, n):
    slide = _base(prs, pal, s.get("title", ""), n)
    steps = (s.get("steps") or [])[:6]
    k = max(1, len(steps))
    y = 4.0
    _rect(slide, 0.8, y - 0.03, SW - 1.6, 0.06, pal["muted"])
    seg = (SW - 1.6) / k
    for i, st in enumerate(steps):
        cx = 0.8 + seg * i + seg / 2
        _rect(slide, cx - 0.18, y - 0.18, 0.36, 0.36, pal["accent"], MSO_SHAPE.OVAL)
        up = i % 2 == 0
        ty = y - 1.95 if up else y + 0.4
        _text(slide, cx - seg / 2 + 0.05, ty, seg - 0.1, 0.5, st.get("title", ""), 20, pal["accent"], True, PP_ALIGN.CENTER)
        _text(slide, cx - seg / 2 + 0.05, ty + 0.5, seg - 0.1, 1.2, st.get("text", ""), 15, pal["text"], align=PP_ALIGN.CENTER)
    return slide


def _slide_stats(prs, pal, s, n):
    slide = _base(prs, pal, s.get("title", ""), n)
    stats = (s.get("stats") or [])[:4]
    k = max(1, len(stats))
    gap = 0.4
    w = (SW - 1.2 - gap * (k - 1)) / k
    for i, st in enumerate(stats):
        x = 0.6 + i * (w + gap)
        _rect(slide, x, 2.2, w, 3.2, pal["surface"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        val = str(st.get("value", ""))
        _text(slide, x, 2.8, w, 1.3, val, 54 if len(val) <= 6 else 38, pal["accent"] if i % 2 == 0 else pal["accent2"],
              True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        _text(slide, x + 0.2, 4.2, w - 0.4, 1.0, st.get("label", ""), 20, pal["muted"], align=PP_ALIGN.CENTER)
    if s.get("bullets"):
        _bullets(slide, 0.9, 5.7, SW - 1.8, 1.2, s.get("bullets"), pal, 16)
    return slide


def _slide_quote(prs, pal, s, n):
    slide = _base(prs, pal, None, n)
    _text(slide, 1.0, 1.0, 2, 2, "“", 140, pal["accent"], True)
    q = str(s.get("quote") or s.get("title", ""))
    _text(slide, 1.6, 2.3, SW - 3.2, 3, q, 34 if len(q) < 120 else 26, pal["text"], False, italic=True)
    if s.get("author"):
        _text(slide, 1.6, 5.5, SW - 3.2, 0.6, "— " + str(s["author"]), 20, pal["muted"])
    return slide


def _slide_section(prs, pal, s, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _c(pal["surface"])
    _rect(slide, 0, 0, 0.35, SH, pal["accent"])
    _text(slide, 1.2, 2.7, SW - 2.4, 1.6, s.get("title", ""), 48, pal["text"], True, anchor=MSO_ANCHOR.MIDDLE)
    if s.get("bullets"):
        _text(slide, 1.2, 4.4, SW - 2.4, 1.0, " · ".join(map(str, s["bullets"][:3])), 18, pal["muted"])
    return slide


_LAYOUTS = {
    "bullets": _slide_bullets, "two_column": _slide_two, "comparison": _slide_compare,
    "process": _slide_process, "timeline": _slide_timeline, "stats": _slide_stats,
    "quote": _slide_quote, "section": _slide_section,
}


def build_deck(plan: dict, style: str, path: Path, closing: str = "Thank you") -> int:
    pal = PALETTES.get(str(style).lower(), PALETTES["midnight"])
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    _title_slide(prs, pal, plan.get("title", "Presentation"), plan.get("subtitle", ""), plan.get("notes", ""))
    count = 1
    for s in plan.get("slides") or []:
        if not isinstance(s, dict):
            continue
        fn = _LAYOUTS.get(str(s.get("layout", "bullets")).lower(), _slide_bullets)
        if fn is _slide_bullets and not s.get("bullets") and s.get("steps"):
            fn = _slide_process
        count += 1
        slide = fn(prs, pal, s, count)
        _notes(slide, s.get("notes"))
    _closing_slide(prs, pal, closing)
    prs.save(str(path))
    return count + 1


def create_presentation(parameters: dict, player=None, speak=None) -> str:
    if not _PPTX:
        return need("pptx", "python-pptx")
    p = parameters or {}
    topic = str(p.get("topic") or "").strip()
    style = str(p.get("style") or "midnight").lower()
    if style not in PALETTES:
        style = "midnight"
    raw = p.get("slides_json")
    plan = None
    if raw:
        try:
            plan = json.loads(raw) if isinstance(raw, str) else raw
        except Exception:
            plan = None
    if plan is None:
        if not topic:
            return "Sir, tell me the topic of the presentation."
        n = int(p.get("slide_count") or 8)
        n = max(3, min(20, n))
        lang = str(p.get("language") or "the same language as the topic")
        audience = str(p.get("audience") or "a general audience")
        log(f"[PPT] Planning {n} slides about: {topic}", player)
        plan = ask_json(
            f"Topic: {topic}\nBody slides: {n - 2}\nAudience: {audience}\nLanguage: {lang}\n"
            f"Extra instructions: {p.get('instructions') or 'none'}",
            SYSTEM, default=None)
    if not isinstance(plan, dict) or not plan.get("slides"):
        return "Sir, I couldn't plan the slides right now. Please try again, or give me the slide content."
    plan.setdefault("title", topic or "Presentation")
    folder = output_dir("Presentations")
    path = unique_path(folder, safe_name(p.get("filename") or plan["title"], "presentation"), ".pptx")
    try:
        total = build_deck(plan, style, path, str(p.get("closing") or "Thank you"))
    except Exception as e:
        return f"Sir, building the presentation failed: {e}"
    log(f"[PPT] Saved {path}", player)
    if p.get("open_after", True):
        open_path(path)
    return f"Presentation ready, sir: {total} slides with speaker notes, style '{style}'. Saved to {path}."


TOOL = {
    "name": "create_presentation",
    "description": (
        "Builds a complete PowerPoint (.pptx) deck from a topic: title slide, structured body slides "
        "(bullets, two-column, comparison, process and timeline diagrams, stats, quotes), speaker notes and "
        "a closing slide. Use when the user asks for a presentation, slides, a deck or a PowerPoint."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "topic": {"type": "STRING", "description": "What the presentation is about"},
            "slide_count": {"type": "INTEGER", "description": "Total slides including title and closing (3-20, default 8)"},
            "style": {"type": "STRING", "description": "midnight | ocean | sunset | forest | mono | aurora"},
            "audience": {"type": "STRING", "description": "Who it is for"},
            "language": {"type": "STRING", "description": "Language of the slides"},
            "instructions": {"type": "STRING", "description": "Any extra requirements"},
            "filename": {"type": "STRING", "description": "Optional file name"},
            "open_after": {"type": "BOOLEAN", "description": "Open the file when done (default true)"},
        },
        "required": ["topic"],
    },
    "handler": create_presentation,
}

# by ASDcuber521
